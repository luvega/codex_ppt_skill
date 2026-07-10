from __future__ import annotations

import json
import importlib.util
import hashlib
import subprocess
import sys
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Inches, Pt


SKILL_ROOT = Path(__file__).resolve().parents[1]
SCRIPTS = SKILL_ROOT / "scripts"
ROOT = SKILL_ROOT.parents[2]
GENERATOR_SPEC = importlib.util.spec_from_file_location(
    "generate_style_discovery_previews",
    SCRIPTS / "generate_style_discovery_previews.py",
)
assert GENERATOR_SPEC and GENERATOR_SPEC.loader
GENERATOR = importlib.util.module_from_spec(GENERATOR_SPEC)
GENERATOR_SPEC.loader.exec_module(GENERATOR)


def run_script(name: str, *args: object) -> subprocess.CompletedProcess[str]:
    return subprocess.run(
        [sys.executable, str(SCRIPTS / name), *(str(value) for value in args)],
        cwd=ROOT,
        text=True,
        capture_output=True,
        encoding="utf-8",
    )


def write_brief(path: Path, **overrides: object) -> dict[str, object]:
    payload: dict[str, object] = {
        "workflow_mode": "new_deck",
        "purpose": "科研报告",
        "audience": "教师与研究生",
        "duration_minutes": 15,
        "target_slide_count": 12,
        "content_readiness": "complete",
        "delivery_mode": "speaker_led",
        "brand_fidelity": "strict",
        "requested_style_id": None,
        "style_discovery_required": True,
        "content_confirmed": True,
        "title": "单细胞图谱揭示治疗后免疫状态重塑",
        "subtitle": "面向课题组的阶段性科研报告",
        "author": "中山大学研究团队",
        "date": "2026-07-10",
    }
    payload.update(overrides)
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    return payload


def write_asset_review(path: Path, assets: list[dict[str, object]] | None = None) -> Path:
    rows = assets or []
    payload = {"status": "reviewed" if rows else "none_provided", "assets": rows}
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    return path


def run_generator(
    brief_path: Path,
    output_dir: Path,
    *extra: object,
    asset_review_path: Path | None = None,
) -> subprocess.CompletedProcess[str]:
    review_path = asset_review_path or write_asset_review(brief_path.with_name("asset-review.json"))
    return run_script(
        "generate_style_discovery_previews.py",
        "--brief",
        brief_path,
        "--asset-review",
        review_path,
        "--output-dir",
        output_dir,
        "--render-method",
        "pillow",
        *extra,
    )


def test_palette_for_uses_source_color_variant() -> None:
    spec = {
        "source": "templates/source/sysu-official/中山大学幻灯片模板-绿.pptx",
        "palette_or_colors": {
            "blue": {"accent": "111111", "text": "222222"},
            "green": {"accent": "2F6F4E", "text": "1E2A22"},
        },
    }

    assert GENERATOR.palette_for(spec)["accent"] == "2F6F4E"


def test_general_discovery_defaults_to_blue_family() -> None:
    brief = {
        "title": "课程反馈提升学习成效",
        "subtitle": "通用课程汇报",
        "purpose": "面向教师的课程汇报",
        "audience": "高校教师",
        "delivery_mode": "speaker_led",
    }

    selected = GENERATOR.select_styles(brief)
    identifiers = [entry["id"] for _role, (entry, _spec, _score) in selected]

    assert identifiers[:2] == ["strict-sysu-official-blue", "beamer-sysu-blue"]


def test_generate_style_discovery_outputs_real_content(tmp_path: Path) -> None:
    brief_path = tmp_path / "deck-brief.json"
    brief = write_brief(brief_path)
    output_dir = tmp_path / "style-discovery"

    result = run_generator(
        brief_path,
        output_dir,
        "--select",
        "B",
        "--showcase",
    )

    assert result.returncode == 0, result.stdout + result.stderr
    pptx_path = output_dir / "style-options.pptx"
    assert pptx_path.exists()
    assert (output_dir / "style-discovery-showcase.pptx").exists()
    prs = Presentation(pptx_path)
    assert len(prs.slides) == 3
    assert round(prs.slide_width / 914400, 3) == 13.333
    assert round(prs.slide_height / 914400, 3) == 7.5

    forbidden = ("Option A", "Option B", "Option C", "preview", "template", "style option")
    layout_signatures: list[tuple[tuple[int, int, int, int], ...]] = []
    for slide in prs.slides:
        visible = "\n".join(shape.text for shape in slide.shapes if getattr(shape, "has_text_frame", False))
        assert brief["title"] in visible
        assert not any(token.casefold() in visible.casefold() for token in forbidden)
        layout_signatures.append(
            tuple((shape.left, shape.top, shape.width, shape.height) for shape in slide.shapes)
        )
        for shape in slide.shapes:
            if not getattr(shape, "has_text_frame", False):
                continue
            for paragraph in shape.text_frame.paragraphs:
                for run in paragraph.runs:
                    assert not str(run.font.name or "").startswith("[")
    assert len(set(layout_signatures)) == 3

    for name in ("option-a.png", "option-b.png", "option-c.png"):
        with Image.open(output_dir / name) as image:
            assert image.size == (1600, 900)
    with Image.open(output_dir / "contact-sheet.png") as contact_sheet:
        assert contact_sheet.size == (1600, 620)

    selection = json.loads((output_dir / "style-selection.json").read_text(encoding="utf-8"))
    assert [item["role"] for item in selection["options"]] == ["safe", "structured", "exploratory"]
    assert len({item["style_id"] for item in selection["options"]}) == 3
    assert all(item["name"] for item in selection["options"])
    assert selection["selected_option"] == "B"
    assert selection["confirmed"] is True
    assert selection["approval_scope"] == "none"
    assert selection["render_method"] == "pillow"
    assert selection["input_hashes"]["brief_sha256"]
    assert selection["input_hashes"]["asset_review_sha256"]


def test_approved_asset_is_inserted_into_every_preview(tmp_path: Path) -> None:
    brief_path = tmp_path / "deck-brief.json"
    write_brief(brief_path)
    asset_path = tmp_path / "approved-logo.png"
    Image.new("RGB", (320, 120), "#2F6F4E").save(asset_path)
    review_path = write_asset_review(
        tmp_path / "asset-review.json",
        [
            {
                "path": str(asset_path),
                "kind": "logo",
                "content_summary": "Approved project logo",
                "usable": True,
                "reason": "Provided for cover previews",
                "dominant_colors": ["#2F6F4E"],
                "aspect_ratio": 2.667,
                "source": "user",
                "planned_slide_roles": ["cover"],
                "processing": [],
            }
        ],
    )
    output_dir = tmp_path / "style-discovery"

    generated = run_generator(
        brief_path,
        output_dir,
        asset_review_path=review_path,
    )

    assert generated.returncode == 0, generated.stdout + generated.stderr
    prs = Presentation(output_dir / "style-options.pptx")
    assert all(
        any(shape.shape_type == MSO_SHAPE_TYPE.PICTURE for shape in slide.shapes)
        for slide in prs.slides
    )
    selection = json.loads((output_dir / "style-selection.json").read_text(encoding="utf-8"))
    assert selection["approved_asset_count"] == 1


def test_style_options_pptx_is_byte_reproducible(tmp_path: Path) -> None:
    brief_path = tmp_path / "deck-brief.json"
    write_brief(brief_path)
    review_path = write_asset_review(tmp_path / "asset-review.json")
    first = tmp_path / "first"
    second = tmp_path / "second"

    first_result = run_generator(brief_path, first, asset_review_path=review_path)
    second_result = run_generator(brief_path, second, asset_review_path=review_path)

    assert first_result.returncode == 0, first_result.stdout + first_result.stderr
    assert second_result.returncode == 0, second_result.stdout + second_result.stderr
    first_hash = hashlib.sha256((first / "style-options.pptx").read_bytes()).hexdigest()
    second_hash = hashlib.sha256((second / "style-options.pptx").read_bytes()).hexdigest()
    assert first_hash == second_hash


def test_validate_style_discovery_rejects_internal_labels(tmp_path: Path) -> None:
    brief_path = tmp_path / "deck-brief.json"
    write_brief(brief_path)
    output_dir = tmp_path / "style-discovery"
    generated = run_generator(brief_path, output_dir)
    assert generated.returncode == 0, generated.stdout + generated.stderr

    prs = Presentation(output_dir / "style-options.pptx")
    box = prs.slides[0].shapes.add_textbox(Inches(1), Inches(1), Inches(2), Inches(0.4))
    box.text = "Option A preview"
    prs.save(output_dir / "style-options.pptx")

    result = run_script(
        "validate_style_discovery.py",
        output_dir,
        "--brief",
        brief_path,
    )
    assert result.returncode == 1
    assert "internal preview label" in result.stdout


def test_validate_style_discovery_accepts_valid_outputs(tmp_path: Path) -> None:
    brief_path = tmp_path / "deck-brief.json"
    write_brief(brief_path)
    output_dir = tmp_path / "style-discovery"
    generated = run_generator(brief_path, output_dir, "--select", "A")
    assert generated.returncode == 0, generated.stdout + generated.stderr

    result = run_script(
        "validate_style_discovery.py",
        output_dir,
        "--brief",
        brief_path,
        "--require-confirmed",
    )
    assert result.returncode == 0, result.stdout + result.stderr
    assert "0 error(s)" in result.stdout


def test_validate_style_discovery_enforces_role_groups(tmp_path: Path) -> None:
    brief_path = tmp_path / "deck-brief.json"
    write_brief(brief_path)
    output_dir = tmp_path / "style-discovery"
    generated = run_generator(brief_path, output_dir)
    assert generated.returncode == 0, generated.stdout + generated.stderr
    selection_path = output_dir / "style-selection.json"
    selection = json.loads(selection_path.read_text(encoding="utf-8"))
    selection["options"][0].update(
        {
            "style_id": "beamer-sysu-blue",
            "name": "Beamer SYSU Blue",
            "generation_status": "ready",
        }
    )
    selection_path.write_text(json.dumps(selection, ensure_ascii=False), encoding="utf-8")

    result = run_script(
        "validate_style_discovery.py",
        output_dir,
        "--brief",
        brief_path,
    )

    assert result.returncode == 1
    assert "safe option must use a ready strict-original style" in result.stdout


def test_validator_rejects_nonobject_selection_and_duplicate_letters(tmp_path: Path) -> None:
    brief_path = tmp_path / "deck-brief.json"
    write_brief(brief_path)
    output_dir = tmp_path / "style-discovery"
    assert run_generator(brief_path, output_dir).returncode == 0
    selection_path = output_dir / "style-selection.json"
    original = json.loads(selection_path.read_text(encoding="utf-8"))

    selection_path.write_text("[]", encoding="utf-8")
    nonobject = run_script("validate_style_discovery.py", output_dir, "--brief", brief_path)
    assert nonobject.returncode == 1
    assert "must be a JSON object" in nonobject.stdout
    assert "Traceback" not in nonobject.stderr

    for option in original["options"]:
        option["option"] = "A"
    selection_path.write_text(json.dumps(original, ensure_ascii=False), encoding="utf-8")
    duplicate = run_script("validate_style_discovery.py", output_dir, "--brief", brief_path)
    assert duplicate.returncode == 1
    assert "option letters must be A, B, C in order" in duplicate.stdout


def test_validator_rejects_blank_png_and_unrendered_requirement(tmp_path: Path) -> None:
    brief_path = tmp_path / "deck-brief.json"
    write_brief(brief_path)
    output_dir = tmp_path / "style-discovery"
    assert run_generator(brief_path, output_dir).returncode == 0

    Image.new("RGB", (1600, 900), "white").save(output_dir / "option-a.png")
    blank = run_script("validate_style_discovery.py", output_dir, "--brief", brief_path)
    assert blank.returncode == 1
    assert "option-a.png is visually blank" in blank.stdout

    assert run_generator(brief_path, output_dir).returncode == 0
    unrendered = run_script(
        "validate_style_discovery.py",
        output_dir,
        "--brief",
        brief_path,
        "--require-powerpoint-rendered",
    )
    assert unrendered.returncode == 1
    assert "PowerPoint-rendered previews are required" in unrendered.stdout


def test_generator_rejects_unconfirmed_redesign(tmp_path: Path) -> None:
    brief_path = tmp_path / "deck-brief.json"
    write_brief(
        brief_path,
        workflow_mode="pptx_redesign",
        content_confirmed=False,
    )

    result = run_generator(brief_path, tmp_path / "style-discovery")

    assert result.returncode == 1
    assert "content_confirmed must be true before style discovery" in result.stdout


def test_generator_rejects_when_discovery_is_not_required(tmp_path: Path) -> None:
    brief_path = tmp_path / "deck-brief.json"
    write_brief(
        brief_path,
        requested_style_id="strict-sysu-official-blue",
        style_discovery_required=False,
    )

    result = run_generator(brief_path, tmp_path / "style-discovery")

    assert result.returncode == 1
    assert "style_discovery_required must be true" in result.stdout


def test_reading_first_excludes_speaker_only_candidate() -> None:
    brief = {
        "title": "Minimal spacious focused methods report",
        "subtitle": "Async technical review",
        "purpose": "Detailed internal review",
        "audience": "Research team",
        "delivery_mode": "reading_first",
    }

    selected = GENERATOR.select_styles(brief)
    exploratory = selected[2][1][0]

    assert exploratory["id"] != "moloch-sysu-minimal"
    assert "reading_first" in exploratory["selection_profile"]["delivery_modes"]


def make_audit_fixture(path: Path) -> None:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    prs.slides.add_slide(prs.slide_layouts[6])
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(10), Inches(4))
    frame = box.text_frame
    frame.clear()
    for index, text in enumerate(("结论一", "结论二", "结论三", "结论四")):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = text * 8
        paragraph.font.size = Pt(18)
    prs.save(path)


def make_valid_speaker_fixture(path: Path) -> None:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    cover = prs.slides.add_slide(prs.slide_layouts[6])
    for index, text in enumerate(("真实标题", "副标题", "作者", "日期")):
        box = cover.shapes.add_textbox(Inches(1), Inches(1.3 + index * 0.7), Inches(8), Inches(0.5))
        box.text = text
        box.text_frame.paragraphs[0].font.size = Pt(18)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    title = slide.shapes.add_textbox(Inches(1), Inches(0.35), Inches(10), Inches(0.6))
    title.text_frame.paragraphs[0].text = "标题不属于支撑段落"
    title.text_frame.paragraphs[0].font.size = Pt(24)
    body = slide.shapes.add_textbox(Inches(1), Inches(1.5), Inches(10), Inches(3.5))
    body.text_frame.clear()
    for index, text in enumerate(("支撑点一", "支撑点二", "支撑点三")):
        paragraph = body.text_frame.paragraphs[0] if index == 0 else body.text_frame.add_paragraph()
        paragraph.text = text
        paragraph.font.size = Pt(18)
    prs.save(path)


def test_taste_audit_uses_delivery_mode_brief(tmp_path: Path) -> None:
    pptx = tmp_path / "deck.pptx"
    make_audit_fixture(pptx)
    style = {
        "taste_profile": {
            "mode": "evolve",
            "layout_variance": 4,
            "visual_density": 5,
            "visual_energy": 4,
            "shape_system": "subtle",
            "layout_repetition_limit": 2,
        },
        "layout_pattern_ids": ["claim-summary"],
    }
    style_path = tmp_path / "style.json"
    style_path.write_text(json.dumps(style), encoding="utf-8")
    brief_path = tmp_path / "deck-brief.json"
    write_brief(brief_path, delivery_mode="speaker_led")

    result = run_script(
        "audit_deck_taste.py",
        pptx,
        "--style",
        style_path,
        "--brief",
        brief_path,
    )

    assert result.returncode == 1
    assert "speaker_led" in result.stdout
    assert "more than 3 supporting paragraphs" in result.stdout


def test_taste_audit_does_not_count_title_as_supporting_paragraph(tmp_path: Path) -> None:
    pptx = tmp_path / "deck.pptx"
    make_valid_speaker_fixture(pptx)
    style = {
        "taste_profile": {
            "mode": "evolve",
            "layout_variance": 4,
            "visual_density": 5,
            "visual_energy": 4,
            "shape_system": "subtle",
            "layout_repetition_limit": 2,
        },
        "layout_pattern_ids": ["claim-summary"],
    }
    style_path = tmp_path / "style.json"
    style_path.write_text(json.dumps(style), encoding="utf-8")
    brief_path = tmp_path / "deck-brief.json"
    write_brief(brief_path, delivery_mode="speaker_led")

    result = run_script(
        "audit_deck_taste.py",
        pptx,
        "--style",
        style_path,
        "--brief",
        brief_path,
    )

    assert result.returncode == 0, result.stdout + result.stderr
