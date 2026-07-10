from __future__ import annotations

import base64
import hashlib
import json
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches


SCRIPT = (
    Path(__file__).resolve().parents[1] / "scripts" / "extract_deck_content.py"
)

# A tiny valid PNG. Keeping the fixture self-contained avoids a test dependency on Pillow.
PNG_BYTES = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVQIHWP4z8DwHwAFgAI/"
    "n5ZV6QAAAABJRU5ErkJggg=="
)


def create_fixture(path: Path) -> dict[str, object]:
    image_path = path.with_suffix(".png")
    image_path.write_bytes(PNG_BYTES)

    presentation = Presentation()
    first = presentation.slides.add_slide(presentation.slide_layouts[6])
    title = first.shapes.add_textbox(Inches(0.5), Inches(0.4), Inches(5), Inches(0.6))
    title.text = "研究报告标题"
    body = first.shapes.add_textbox(Inches(0.75), Inches(1.5), Inches(6), Inches(1.25))
    body.text = "第一张幻灯片正文\n保留中文内容"
    picture = first.shapes.add_picture(
        str(image_path), Inches(7), Inches(1.25), width=Inches(1.5), height=Inches(1.5)
    )
    notes_text = "讲者备注：说明研究背景"
    first.notes_slide.notes_text_frame.text = notes_text

    second = presentation.slides.add_slide(presentation.slide_layouts[6])
    second_title = second.shapes.add_textbox(
        Inches(0.5), Inches(0.4), Inches(5), Inches(0.6)
    )
    second_title.text = "第二页结论"
    table = second.shapes.add_table(2, 2, Inches(0.75), Inches(1.4), Inches(6), Inches(2)).table
    table.cell(0, 0).text = "指标"
    table.cell(0, 1).text = "结果"
    table.cell(1, 0).text = "完成率"
    table.cell(1, 1).text = "68%"

    presentation.save(path)
    return {
        "title_geometry": (title.left, title.top, title.width, title.height),
        "picture_geometry": (picture.left, picture.top, picture.width, picture.height),
        "notes_text": notes_text,
    }


class ExtractDeckContentCliTests(unittest.TestCase):
    def test_extracts_ordered_chinese_text_notes_and_picture_metadata(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            workspace = Path(temp_dir)
            input_path = workspace / "fixture.pptx"
            expected = create_fixture(input_path)
            initial_hash = hashlib.sha256(input_path.read_bytes()).hexdigest()
            output_dir = workspace / "extract"
            stale_dir = output_dir / "extracted-assets"
            stale_dir.mkdir(parents=True)
            stale_asset = stale_dir / "stale.png"
            stale_asset.write_bytes(b"stale")

            result = subprocess.run(
                [sys.executable, str(SCRIPT), str(input_path), "--output-dir", str(output_dir)],
                text=True,
                capture_output=True,
                check=False,
            )

            self.assertEqual(result.returncode, 0, result.stderr)
            self.assertEqual(
                hashlib.sha256(input_path.read_bytes()).hexdigest(), initial_hash
            )

            report_path = output_dir / "source-deck-extract.json"
            self.assertTrue(report_path.is_file())
            self.assertFalse(stale_asset.exists())
            report = json.loads(report_path.read_text(encoding="utf-8"))
            self.assertEqual(report["source_path"], str(input_path.resolve()).replace("\\", "/"))
            self.assertEqual(report["sha256_before_extraction"], initial_hash)
            self.assertEqual(report["sha256_after_extraction"], initial_hash)
            self.assertTrue(report["source_unchanged"])
            self.assertGreater(report["slide_size"]["width_emu"], 0)
            self.assertGreater(report["slide_size"]["height_emu"], 0)
            self.assertEqual(report["slide_order"], [0, 1])

            slides = report["slides"]
            self.assertEqual([slide["index"] for slide in slides], [0, 1])
            self.assertEqual([slide["title"] for slide in slides], ["研究报告标题", "第二页结论"])
            self.assertEqual(
                slides[1]["tables"][0]["cells"],
                [["指标", "结果"], ["完成率", "68%"]],
            )

            first = slides[0]
            self.assertEqual(first["speaker_notes"], expected["notes_text"])
            text_by_value = {shape["text"]: shape for shape in first["text_shapes"]}
            self.assertIn("第一张幻灯片正文\n保留中文内容", text_by_value)
            self.assertEqual(
                text_by_value["研究报告标题"]["geometry"],
                {
                    "left_emu": expected["title_geometry"][0],
                    "top_emu": expected["title_geometry"][1],
                    "width_emu": expected["title_geometry"][2],
                    "height_emu": expected["title_geometry"][3],
                },
            )

            self.assertEqual(len(first["pictures"]), 1)
            extracted_picture = first["pictures"][0]
            self.assertEqual(extracted_picture["original_extension"], ".png")
            self.assertEqual(
                extracted_picture["geometry"],
                {
                    "left_emu": expected["picture_geometry"][0],
                    "top_emu": expected["picture_geometry"][1],
                    "width_emu": expected["picture_geometry"][2],
                    "height_emu": expected["picture_geometry"][3],
                },
            )
            self.assertEqual(
                extracted_picture["crop"],
                {"left": 0.0, "top": 0.0, "right": 0.0, "bottom": 0.0},
            )
            self.assertRegex(
                extracted_picture["output_path"],
                r"^extracted-assets/slide-001-picture-001-[0-9a-f]{12}\.png$",
            )
            extracted_path = output_dir / extracted_picture["output_path"]
            self.assertEqual(extracted_path.read_bytes(), PNG_BYTES)

    def test_rejects_a_missing_input_file(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            workspace = Path(temp_dir)
            missing_path = workspace / "does-not-exist.pptx"

            result = subprocess.run(
                [
                    sys.executable,
                    str(SCRIPT),
                    str(missing_path),
                    "--output-dir",
                    str(workspace / "extract"),
                ],
                text=True,
                capture_output=True,
                check=False,
            )

            self.assertNotEqual(result.returncode, 0)
            self.assertIn("Input PPTX does not exist", result.stderr)


if __name__ == "__main__":
    unittest.main()
