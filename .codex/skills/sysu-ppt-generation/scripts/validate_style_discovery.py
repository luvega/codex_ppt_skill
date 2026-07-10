#!/usr/bin/env python
"""Validate the file and metadata contract for style-discovery outputs."""

from __future__ import annotations

import argparse
import json
from pathlib import Path
from typing import Any

from PIL import Image
from pptx import Presentation

from generate_style_discovery_previews import (
    OPTIONS,
    REGISTRY_PATH,
    ROLES,
    load_json,
    sha256_file,
    validate_asset_review,
    validate_brief,
)


def visible_text(slide) -> str:
    return "\n".join(shape.text for shape in slide.shapes if getattr(shape, "has_text_frame", False))


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("output_dir", type=Path)
    parser.add_argument("--brief", type=Path, required=True)
    parser.add_argument("--asset-review", type=Path)
    parser.add_argument("--require-confirmed", action="store_true")
    parser.add_argument("--require-powerpoint-rendered", action="store_true")
    args = parser.parse_args()
    errors: list[str] = []
    try:
        brief = load_json(args.brief)
        errors.extend(validate_brief(brief))
    except (OSError, ValueError, json.JSONDecodeError) as error:
        errors.append(f"invalid deck brief: {error}")
        brief = {}

    asset_review_path = args.asset_review or args.brief.with_name("asset-review.json")
    try:
        asset_review = load_json(asset_review_path)
        asset_errors, approved_assets = validate_asset_review(asset_review, asset_review_path)
        errors.extend(asset_errors)
    except (OSError, ValueError, json.JSONDecodeError) as error:
        errors.append(f"invalid asset review: {error}")
        approved_assets = []

    required = ["style-options.pptx", "option-a.png", "option-b.png", "option-c.png", "contact-sheet.png", "style-selection.json"]
    for name in required:
        if not (args.output_dir / name).is_file():
            errors.append(f"missing required file: {name}")
    selection: dict[str, Any] = {}
    selection_path = args.output_dir / "style-selection.json"
    if selection_path.is_file():
        try:
            loaded_selection = load_json(selection_path)
            if not isinstance(loaded_selection, dict):
                errors.append("style-selection.json must be a JSON object")
            else:
                selection = loaded_selection
        except (OSError, ValueError, json.JSONDecodeError) as error:
            errors.append(f"invalid style-selection.json: {error}")

    registry = {item.get("id"): item for item in load_json(REGISTRY_PATH).get("styles", []) if isinstance(item, dict)}
    options = selection.get("options", []) if isinstance(selection, dict) else []
    if not isinstance(options, list) or len(options) != 3:
        errors.append("style-selection.json must contain three options")
    else:
        roles = [item.get("role") for item in options if isinstance(item, dict)]
        letters = [item.get("option") for item in options if isinstance(item, dict)]
        identifiers = [item.get("style_id") for item in options if isinstance(item, dict)]
        if roles != list(ROLES):
            errors.append("style roles must be safe, structured, exploratory")
        if letters != list(OPTIONS):
            errors.append("style option letters must be A, B, C in order")
        if len(set(identifiers)) != 3:
            errors.append("style option IDs must be distinct")
        for item in options:
            style_id = item.get("style_id") if isinstance(item, dict) else None
            entry = registry.get(style_id)
            if not entry:
                errors.append(f"unregistered style ID: {style_id}")
            elif item.get("generation_status") != entry.get("generation_status"):
                errors.append(f"style status mismatch for {style_id}")
            if not isinstance(item.get("name"), str) or not item.get("name", "").strip():
                errors.append(f"style option {style_id} must include a display name")
        role_entries = {
            item.get("role"): registry.get(item.get("style_id"))
            for item in options
            if isinstance(item, dict)
        }
        safe = role_entries.get("safe") or {}
        if safe.get("group") != "strict-original" or safe.get("generation_status") != "ready":
            errors.append("safe option must use a ready strict-original style")
        structured = role_entries.get("structured") or {}
        if structured.get("group") != "beamer-inspired" or structured.get("generation_status") != "ready":
            errors.append("structured option must use a ready beamer-inspired style")
        exploratory = role_entries.get("exploratory") or {}
        if exploratory.get("generation_status") not in {"ready", "style_selection_only"}:
            errors.append("exploratory option must use a registered selectable style")

    chosen = selection.get("selected_option") if isinstance(selection, dict) else None
    if chosen not in (*OPTIONS, None):
        errors.append("selected_option must be A, B, C, or null")
    selected_item = next((item for item in options if isinstance(item, dict) and item.get("option") == chosen), None)
    if chosen and not selected_item:
        errors.append("selected option is not present in options")
    if selected_item and selection.get("selected_style_id") != selected_item.get("style_id"):
        errors.append("selected_style_id does not match the selected option")
    if args.require_confirmed and (not selection.get("confirmed") or chosen not in OPTIONS):
        errors.append("style selection confirmation is required")
    if selected_item and selected_item.get("generation_status") == "style_selection_only" and selection.get("approval_scope") != "deck_only":
        errors.append("selection-only style requires approval_scope deck_only")
    if selected_item and selected_item.get("generation_status") == "ready" and selection.get("approval_scope") != "none":
        errors.append("ready style selection must use approval_scope none")
    render_method = selection.get("render_method")
    if render_method not in {"powerpoint", "pillow"}:
        errors.append("style selection must record render_method")
    if args.require_powerpoint_rendered and render_method != "powerpoint":
        errors.append("PowerPoint-rendered previews are required")
    if selection.get("approved_asset_count") != len(approved_assets):
        errors.append("approved_asset_count does not match asset review")
    hashes = selection.get("input_hashes")
    if not isinstance(hashes, dict):
        errors.append("style selection must record input_hashes")
    else:
        hash_paths = {
            "brief_sha256": args.brief,
            "asset_review_sha256": asset_review_path,
            "style_registry_sha256": REGISTRY_PATH,
            "generator_sha256": Path(__file__).with_name("generate_style_discovery_previews.py"),
        }
        for key, path in hash_paths.items():
            if path.is_file() and hashes.get(key) != sha256_file(path):
                errors.append(f"input hash mismatch: {key}")
        style_hashes = hashes.get("style_specs_sha256")
        if not isinstance(style_hashes, dict):
            errors.append("input hashes must include style_specs_sha256")
        else:
            for item in options:
                if not isinstance(item, dict):
                    continue
                entry = registry.get(item.get("style_id"))
                if not entry:
                    continue
                spec_path = REGISTRY_PATH.parents[2] / entry["style_spec"]
                if spec_path.is_file() and style_hashes.get(entry["id"]) != sha256_file(spec_path):
                    errors.append(f"input hash mismatch: style spec {entry['id']}")

    pptx_path = args.output_dir / "style-options.pptx"
    if pptx_path.is_file():
        try:
            prs = Presentation(pptx_path)
            if len(prs.slides) != 3:
                errors.append("style-options.pptx must contain three slides")
            if round(prs.slide_width / 914400, 3) != 13.333 or round(prs.slide_height / 914400, 3) != 7.5:
                errors.append("style-options.pptx must be 13.333x7.5 inches")
            forbidden = {"option a", "option b", "option c", "style option", "preview", "template"}
            forbidden.update(str(item.get("style_id", "")).casefold() for item in options if isinstance(item, dict))
            forbidden.update(str(brief.get(key, "")).casefold() for key in ("workflow_mode", "delivery_mode", "content_readiness"))
            for number, slide in enumerate(prs.slides, 1):
                text = visible_text(slide)
                if brief.get("title") and brief["title"] not in text:
                    errors.append(f"slide {number}: deck title is missing")
                if any(token and token in text.casefold() for token in forbidden):
                    errors.append(f"slide {number}: internal preview label leakage")
        except Exception as error:  # python-pptx normalizes several ZIP/XML errors.
            errors.append(f"cannot read style-options.pptx: {error}")

    for name in ("option-a.png", "option-b.png", "option-c.png"):
        path = args.output_dir / name
        if path.is_file():
            try:
                with Image.open(path) as image:
                    if image.size != (1600, 900):
                        errors.append(f"{name} must be 1600x900 pixels")
                    colors = image.convert("RGB").getcolors(maxcolors=2)
                    if colors is not None and len(colors) <= 1:
                        errors.append(f"{name} is visually blank")
                    image.resize((800, 450)).load()
            except OSError as error:
                errors.append(f"cannot read {name}: {error}")

    contact_sheet = args.output_dir / "contact-sheet.png"
    if contact_sheet.is_file():
        try:
            with Image.open(contact_sheet) as image:
                if image.size != (1600, 620):
                    errors.append("contact-sheet.png must be 1600x620 pixels")
        except OSError as error:
            errors.append(f"cannot read contact-sheet.png: {error}")

    for error in errors:
        print(f"ERROR: {error}")
    print(f"Style discovery validation: {len(errors)} error(s)")
    return 1 if errors else 0


if __name__ == "__main__":
    raise SystemExit(main())
