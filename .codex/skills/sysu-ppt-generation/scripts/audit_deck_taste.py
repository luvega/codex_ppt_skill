#!/usr/bin/env python
"""Read-only mechanical taste audit for a PPTX deck."""

from __future__ import annotations

import argparse
import json
import re
from collections import Counter
from pathlib import Path
from typing import Any

from pptx import Presentation


REQUIRED_PROFILE = {"mode", "layout_variance", "visual_density", "visual_energy", "shape_system", "layout_repetition_limit"}
MODES = {"preserve", "evolve", "selection"}


def load_json(path: Path) -> Any:
    return json.loads(path.read_text(encoding="utf-8"))


def iter_runs(shape):
    if not getattr(shape, "has_text_frame", False):
        return
    for paragraph in shape.text_frame.paragraphs:
        for run in paragraph.runs:
            if run.text.strip():
                yield run


def pattern_sequence(mapping: Any) -> list[str | None]:
    rows = mapping.get("slides", mapping) if isinstance(mapping, dict) else mapping
    if isinstance(rows, dict):
        rows = list(rows.values())
    if not isinstance(rows, list):
        return []
    return [row.get("layout_pattern_id") if isinstance(row, dict) else None for row in rows]


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("pptx", type=Path)
    parser.add_argument("--style", type=Path, required=True)
    parser.add_argument("--mapping", type=Path)
    args = parser.parse_args()

    style = load_json(args.style)
    profile = style.get("taste_profile", {})
    errors: list[str] = []
    warnings: list[str] = []
    infos: list[str] = []

    missing = REQUIRED_PROFILE - set(profile)
    if missing:
        errors.append(f"style missing taste profile fields: {', '.join(sorted(missing))}")
    if profile.get("mode") not in MODES:
        errors.append(f"invalid taste mode: {profile.get('mode')}")
    for key in ("layout_variance", "visual_density", "visual_energy"):
        value = profile.get(key)
        if not isinstance(value, int) or not 1 <= value <= 10:
            errors.append(f"{key} must be an integer from 1 to 10")

    prs = Presentation(args.pptx)
    width, height = prs.slide_width, prs.slide_height
    density = profile.get("visual_density", 5)
    density_limit = 160 if density <= 3 else 280 if density <= 6 else 440
    fonts: Counter[str] = Counter()

    for slide_no, slide in enumerate(prs.slides, 1):
        visible = 0
        preserve = profile.get("mode") == "preserve"
        for shape in slide.shapes:
            if shape.left < 0 or shape.top < 0 or shape.left + shape.width > width or shape.top + shape.height > height:
                message = f"slide {slide_no}: shape outside slide bounds"
                (warnings if preserve else errors).append(message)
            if getattr(shape, "has_text_frame", False):
                raw = shape.text.strip()
                visible += len(re.findall(r"[\u3400-\u9fff]|[A-Za-z0-9]+", raw))
                for run in iter_runs(shape):
                    if run.font.name:
                        fonts[run.font.name] += 1
                    if run.font.size:
                        size = run.font.size.pt
                        in_footer = shape.top > height * 0.86
                        substantive = slide_no != 1 and len(run.text.strip()) >= 12 and shape.height > height * 0.05
                        if substantive and not in_footer and size < 16:
                            message = f"slide {slide_no}: substantive text below 16 pt ({size:.1f} pt)"
                            (warnings if preserve else errors).append(message)
        if visible > density_limit * 1.25:
            errors.append(f"slide {slide_no}: text density {visible} exceeds profile budget {density_limit}")
        elif visible > density_limit:
            warnings.append(f"slide {slide_no}: text density {visible} is above profile budget {density_limit}")

    if args.mapping:
        sequence = pattern_sequence(load_json(args.mapping))
        if len(sequence) != len(prs.slides):
            errors.append(f"mapping covers {len(sequence)} slides but PPTX has {len(prs.slides)}")
        allowed = set(style.get("layout_pattern_ids", []))
        limit = profile.get("layout_repetition_limit", 2)
        run_pattern = None; run_length = 0
        for index, pattern in enumerate(sequence, 1):
            if not pattern:
                errors.append(f"slide {index}: missing layout_pattern_id")
                continue
            if pattern not in allowed:
                errors.append(f"slide {index}: unknown or disallowed layout pattern {pattern}")
            run_length = run_length + 1 if pattern == run_pattern else 1
            run_pattern = pattern
            if run_length > limit:
                errors.append(f"slide {index}: layout pattern {pattern} repeated more than {limit} times")

    infos.append(f"slides={len(prs.slides)} size={width/914400:.3f}x{height/914400:.3f}")
    infos.append(f"fonts={', '.join(name for name, _ in fonts.most_common(6)) or 'theme/default'}")
    for item in infos: print(f"INFO: {item}")
    for item in warnings: print(f"WARN: {item}")
    for item in errors: print(f"ERROR: {item}")
    print(f"Taste audit: {len(errors)} error(s), {len(warnings)} warning(s)")
    return 1 if errors else 0


if __name__ == "__main__":
    raise SystemExit(main())
