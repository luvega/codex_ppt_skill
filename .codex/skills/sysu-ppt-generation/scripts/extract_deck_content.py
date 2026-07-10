#!/usr/bin/env python
"""Extract text, speaker notes, and picture assets from a PowerPoint deck."""

from __future__ import annotations

import argparse
import hashlib
import json
import shutil
from pathlib import Path
from typing import Any, Iterable

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.oxml.ns import qn


def geometry(shape: Any) -> dict[str, int]:
    return {
        "left_emu": int(shape.left),
        "top_emu": int(shape.top),
        "width_emu": int(shape.width),
        "height_emu": int(shape.height),
    }


def iter_shapes(shapes: Iterable[Any]) -> Iterable[Any]:
    for shape in shapes:
        yield shape
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from iter_shapes(shape.shapes)


def speaker_notes(slide: Any) -> str:
    try:
        text_frame = slide.notes_slide.notes_text_frame
    except (AttributeError, KeyError):
        return ""
    return text_frame.text if text_frame is not None else ""


def picture_extension(shape: Any) -> str:
    extension = str(shape.image.ext or "bin").lower().lstrip(".")
    return f".{extension}"


def image_relationship_ids(element: Any) -> list[str]:
    return [
        value
        for blip in element.xpath(".//a:blip")
        if (value := blip.get(qn("r:embed")))
    ]


def write_related_image(
    slide: Any,
    relationship_id: str,
    assets_dir: Path,
    slide_number: int,
    image_number: int,
    kind: str,
) -> tuple[str, str]:
    part = slide.part.related_part(relationship_id)
    image_bytes = part.blob
    digest = hashlib.sha256(image_bytes).hexdigest()[:12]
    extension = Path(str(part.partname)).suffix.lower() or ".bin"
    asset_name = f"slide-{slide_number:03d}-{kind}-{image_number:03d}-{digest}{extension}"
    (assets_dir / asset_name).write_bytes(image_bytes)
    return f"extracted-assets/{asset_name}", extension


def extract_deck(source_path: Path, output_dir: Path) -> Path:
    source_path = source_path.resolve()
    source_hash = hashlib.sha256(source_path.read_bytes()).hexdigest()
    presentation = Presentation(str(source_path))

    output_dir = output_dir.resolve()
    output_dir.mkdir(parents=True, exist_ok=True)
    assets_dir = output_dir / "extracted-assets"
    if assets_dir.exists():
        shutil.rmtree(assets_dir)
    assets_dir.mkdir(parents=True, exist_ok=True)
    slides: list[dict[str, Any]] = []

    for slide_index, slide in enumerate(presentation.slides):
        text_shapes: list[dict[str, Any]] = []
        tables: list[dict[str, Any]] = []
        pictures: list[dict[str, Any]] = []
        seen_relationships: set[str] = set()

        for shape_index, shape in enumerate(iter_shapes(slide.shapes), start=1):
            if getattr(shape, "has_text_frame", False):
                text_shapes.append(
                    {
                        "shape_index": shape_index,
                        "name": shape.name,
                        "text": shape.text,
                        "geometry": geometry(shape),
                    }
                )

            if getattr(shape, "has_table", False):
                table = shape.table
                tables.append(
                    {
                        "shape_index": shape_index,
                        "name": shape.name,
                        "rows": len(table.rows),
                        "columns": len(table.columns),
                        "cells": [
                            [table.cell(row, column).text for column in range(len(table.columns))]
                            for row in range(len(table.rows))
                        ],
                        "geometry": geometry(shape),
                    }
                )

            if shape.shape_type != MSO_SHAPE_TYPE.PICTURE:
                continue

            picture_index = len(pictures) + 1
            image_bytes = shape.image.blob
            digest = hashlib.sha256(image_bytes).hexdigest()[:12]
            extension = picture_extension(shape)
            asset_name = (
                f"slide-{slide_index + 1:03d}-picture-{picture_index:03d}-{digest}{extension}"
            )
            asset_path = assets_dir / asset_name
            asset_path.write_bytes(image_bytes)
            pictures.append(
                {
                    "shape_index": shape_index,
                    "name": shape.name,
                    "kind": "picture",
                    "output_path": f"extracted-assets/{asset_name}",
                    "original_extension": extension,
                    "geometry": geometry(shape),
                    "crop": {
                        "left": float(shape.crop_left),
                        "top": float(shape.crop_top),
                        "right": float(shape.crop_right),
                        "bottom": float(shape.crop_bottom),
                    },
                }
            )
            seen_relationships.update(image_relationship_ids(shape.element))

        for shape_index, shape in enumerate(iter_shapes(slide.shapes), start=1):
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                continue
            for relationship_id in image_relationship_ids(shape.element):
                if relationship_id in seen_relationships:
                    continue
                seen_relationships.add(relationship_id)
                output_path, extension = write_related_image(
                    slide,
                    relationship_id,
                    assets_dir,
                    slide_index + 1,
                    len(pictures) + 1,
                    "fill",
                )
                pictures.append(
                    {
                        "shape_index": shape_index,
                        "name": shape.name,
                        "kind": "shape-fill",
                        "output_path": output_path,
                        "original_extension": extension,
                        "geometry": geometry(shape),
                        "crop": None,
                    }
                )

        for relationship_id in image_relationship_ids(slide.element):
            if relationship_id in seen_relationships:
                continue
            seen_relationships.add(relationship_id)
            output_path, extension = write_related_image(
                slide,
                relationship_id,
                assets_dir,
                slide_index + 1,
                len(pictures) + 1,
                "background",
            )
            pictures.append(
                {
                    "shape_index": None,
                    "name": "slide-background",
                    "kind": "background",
                    "output_path": output_path,
                    "original_extension": extension,
                    "geometry": {
                        "left_emu": 0,
                        "top_emu": 0,
                        "width_emu": int(presentation.slide_width),
                        "height_emu": int(presentation.slide_height),
                    },
                    "crop": None,
                }
            )

        title_shape = slide.shapes.title
        title = title_shape.text if title_shape is not None and title_shape.text else ""
        if not title:
            title = next((item["text"] for item in text_shapes if item["text"]), "")
        if not title:
            title = next(
                (cell for table in tables for row in table["cells"] for cell in row if cell),
                "",
            )
        slides.append(
            {
                "index": slide_index,
                "title": title,
                "text_shapes": text_shapes,
                "tables": tables,
                "speaker_notes": speaker_notes(slide),
                "pictures": pictures,
            }
        )

    source_hash_after = hashlib.sha256(source_path.read_bytes()).hexdigest()
    if source_hash_after != source_hash:
        raise RuntimeError(f"Source PPTX changed during extraction: {source_path}")

    payload = {
        "source_path": source_path.as_posix(),
        "sha256_before_extraction": source_hash,
        "sha256_after_extraction": source_hash_after,
        "source_unchanged": True,
        "slide_size": {
            "width_emu": int(presentation.slide_width),
            "height_emu": int(presentation.slide_height),
        },
        "slide_order": [slide["index"] for slide in slides],
        "slides": slides,
    }
    output_dir.mkdir(parents=True, exist_ok=True)
    report_path = output_dir / "source-deck-extract.json"
    report_path.write_text(
        json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    return report_path


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("input_pptx", type=Path, metavar="INPUT.pptx")
    parser.add_argument("--output-dir", type=Path, required=True, metavar="DIR")
    args = parser.parse_args()

    if not args.input_pptx.is_file():
        parser.error(f"Input PPTX does not exist: {args.input_pptx}")

    report_path = extract_deck(args.input_pptx, args.output_dir)
    print(f"Wrote {report_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
