#!/usr/bin/env python
"""Create real-content SYSU style-discovery previews without editing templates."""

from __future__ import annotations

import argparse
import hashlib
import json
import os
import re
import shutil
import subprocess
import sys
import zipfile
from importlib.metadata import version
from datetime import datetime
from pathlib import Path
from typing import Any

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt


SKILL_ROOT = Path(__file__).resolve().parents[1]
ROOT = SKILL_ROOT.parents[2]
REGISTRY_PATH = ROOT / "templates" / "styles" / "style-index.json"
WORKFLOW_MODES = {"new_deck", "pptx_redesign", "pptx_revision"}
DELIVERY_MODES = {"speaker_led", "reading_first"}
CONTENT_READINESS = {"complete", "rough_notes", "topic_only"}
BRAND_FIDELITY = {"strict", "sysu_evolved"}
ROLES = ("safe", "structured", "exploratory")
OPTIONS = ("A", "B", "C")
ASSET_REVIEW_STATUSES = {"none_provided", "reviewed"}
REQUIRED_ASSET_FIELDS = {
    "path",
    "kind",
    "content_summary",
    "usable",
    "reason",
    "dominant_colors",
    "aspect_ratio",
    "source",
    "planned_slide_roles",
    "processing",
}
GENERATOR_VERSION = "0.4"
FIXED_TIMESTAMP = datetime(2000, 1, 1, 0, 0, 0)


def load_json(path: Path) -> Any:
    return json.loads(path.read_text(encoding="utf-8"))


def validate_brief(brief: Any) -> list[str]:
    if not isinstance(brief, dict):
        return ["deck brief must be a JSON object"]
    errors: list[str] = []
    values = {
        "workflow_mode": WORKFLOW_MODES,
        "delivery_mode": DELIVERY_MODES,
        "content_readiness": CONTENT_READINESS,
        "brand_fidelity": BRAND_FIDELITY,
    }
    for field, allowed in values.items():
        value = brief.get(field)
        if value not in allowed:
            errors.append(f"invalid {field}: {value!r}")
    if not isinstance(brief.get("title"), str) or not brief["title"].strip():
        errors.append("deck brief title must be non-empty")
    if not isinstance(brief.get("content_confirmed"), bool):
        errors.append("content_confirmed must be a boolean")
    elif brief.get("content_confirmed") is not True:
        errors.append("content_confirmed must be true before style discovery")
    if not isinstance(brief.get("style_discovery_required"), bool):
        errors.append("style_discovery_required must be a boolean")
    elif brief.get("style_discovery_required") is not True:
        errors.append("style_discovery_required must be true to generate previews")
    requested_style = brief.get("requested_style_id")
    if requested_style is not None and (not isinstance(requested_style, str) or not requested_style.strip()):
        errors.append("requested_style_id must be null or a non-empty string")
    if requested_style and brief.get("style_discovery_required") is True:
        errors.append("style discovery must be false when requested_style_id is set")
    if brief.get("workflow_mode") == "pptx_revision" and brief.get("style_discovery_required") is True:
        errors.append("pptx_revision preserves the existing style and does not require discovery")
    if (
        brief.get("workflow_mode") in {"new_deck", "pptx_redesign"}
        and requested_style is None
        and brief.get("style_discovery_required") is not True
    ):
        errors.append("new_deck/pptx_redesign without a requested style requires discovery")
    for field in ("purpose", "audience"):
        if not isinstance(brief.get(field), str) or not brief.get(field, "").strip():
            errors.append(f"{field} must be non-empty")
    for field in ("duration_minutes", "target_slide_count"):
        value = brief.get(field)
        if not isinstance(value, int) or isinstance(value, bool) or value <= 0:
            errors.append(f"{field} must be a positive integer")
    return errors


def validate_asset_review(review: Any, review_path: Path) -> tuple[list[str], list[Path]]:
    if not isinstance(review, dict):
        return ["asset review must be a JSON object"], []
    errors: list[str] = []
    status = review.get("status")
    assets = review.get("assets")
    if status not in ASSET_REVIEW_STATUSES:
        errors.append(f"invalid asset review status: {status!r}")
    if not isinstance(assets, list):
        return errors + ["asset review assets must be a list"], []
    if status == "none_provided" and assets:
        errors.append("none_provided asset review must have an empty assets list")
    approved: list[Path] = []
    for index, item in enumerate(assets, 1):
        if not isinstance(item, dict):
            errors.append(f"asset {index} must be an object")
            continue
        missing = sorted(REQUIRED_ASSET_FIELDS - set(item))
        if missing:
            errors.append(f"asset {index} missing fields: {', '.join(missing)}")
        if not isinstance(item.get("usable"), bool):
            errors.append(f"asset {index} usable must be a boolean")
        if item.get("usable") is not True:
            continue
        raw_path = item.get("path")
        if not isinstance(raw_path, str) or not raw_path.strip():
            errors.append(f"asset {index} path must be non-empty")
            continue
        asset_path = Path(raw_path)
        if not asset_path.is_absolute():
            asset_path = review_path.parent / asset_path
        asset_path = asset_path.resolve()
        if not asset_path.is_file():
            errors.append(f"approved asset does not exist: {raw_path}")
            continue
        try:
            with Image.open(asset_path) as image:
                image.verify()
        except OSError:
            errors.append(f"approved asset is not a readable image: {raw_path}")
            continue
        approved.append(asset_path)
    return errors, approved


def tokenise(brief: dict[str, Any]) -> set[str]:
    values = [brief.get(key, "") for key in ("title", "subtitle", "purpose", "audience")]
    source = " ".join(map(str, values)).casefold()
    terms = set(re.findall(r"[a-z0-9]+|[\u3400-\u9fff]+", source))
    domain_cues = {
        "biomedical life science public health sustainability research": ("单细胞", "免疫", "治疗", "临床", "生物", "医学", "健康", "细胞", "组学"),
        "technical algorithm code computational data research": ("算法", "代码", "计算", "模型", "数据", "人工智能", "ai"),
        "formal defense institutional university report": ("答辩", "正式", "汇报", "学校", "学院"),
        "teaching lecture course structured": ("教学", "课程", "讲座", "课堂"),
    }
    for keywords, cues in domain_cues.items():
        if any(cue in source for cue in cues):
            terms.update(keywords.split())
    return terms


def brief_color_family(brief: dict[str, Any]) -> str:
    source = " ".join(
        str(brief.get(key, ""))
        for key in ("title", "subtitle", "purpose", "audience")
    ).casefold()
    biomedical = ("生物", "医学", "临床", "免疫", "细胞", "组学", "公共卫生", "健康", "生态", "biomedical", "public health")
    formal = ("政策", "行政", "仪式", "党政", "学院级", "学校级", "官方会议", "policy", "ceremonial", "administrative")
    if any(cue in source for cue in biomedical):
        return "green"
    if any(cue in source for cue in formal):
        return "red"
    return "blue"


def style_score(entry: dict[str, Any], terms: set[str]) -> int:
    profile = entry.get("selection_profile", {})
    searchable = [entry.get("use_case", ""), entry.get("recommended_for", "")]
    if isinstance(profile, dict):
        searchable.extend(str(value) for value in profile.values())
    haystack = " ".join(searchable).casefold()
    return sum(1 for term in terms if term and term in haystack)


def select_styles(brief: dict[str, Any]) -> list[tuple[str, dict[str, Any], dict[str, Any]]]:
    registry = load_json(REGISTRY_PATH).get("styles", [])
    if not isinstance(registry, list):
        raise ValueError("style registry has no styles list")
    terms = tokenise(brief)
    delivery_mode = brief.get("delivery_mode")
    candidates: list[tuple[dict[str, Any], int]] = []
    for entry in registry:
        if not isinstance(entry, dict) or not isinstance(entry.get("id"), str):
            continue
        spec_path = ROOT / str(entry.get("style_spec", ""))
        if not spec_path.is_file():
            continue
        modes = entry.get("selection_profile", {}).get("delivery_modes", [])
        if delivery_mode and delivery_mode not in modes:
            continue
        candidates.append((entry, style_score(entry, terms)))
    if len(candidates) < 3:
        raise ValueError("style registry does not contain three usable styles")

    def choose(predicate, used: set[str]) -> tuple[dict[str, Any], int]:
        matches = [item for item in candidates if item[0]["id"] not in used and predicate(item)]
        if not matches:
            matches = [item for item in candidates if item[0]["id"] not in used]
        if not matches:
            raise ValueError("style registry has no compatible unused style")
        return max(matches, key=lambda item: (item[1], item[0]["id"]))

    used: set[str] = set()
    color_family = brief_color_family(brief)
    safe_id = f"strict-sysu-official-{color_family}"
    structured_id = f"beamer-sysu-{color_family}"
    safe = choose(lambda item: item[0].get("id") == safe_id, used)
    used.add(safe[0]["id"])
    structured = choose(lambda item: item[0].get("id") == structured_id, used)
    used.add(structured[0]["id"])
    exploratory = choose(lambda item: item[0].get("generation_status") == "style_selection_only", used)
    selected: list[tuple[str, tuple[dict[str, Any], dict[str, Any], int]]] = []
    for role, (entry, score) in zip(ROLES, (safe, structured, exploratory), strict=True):
        spec = load_json(ROOT / entry["style_spec"])
        selected.append((role, (entry, spec, score)))
    return selected


def palette_for(spec: dict[str, Any]) -> dict[str, str]:
    palette = spec.get("palette", spec.get("palette_or_colors", {}))
    if isinstance(palette, dict) and "accent" not in palette:
        source = str(spec.get("source", "")).casefold()
        variant = next(
            (
                key
                for cue, key in (("蓝", "blue"), ("绿", "green"), ("红", "red"))
                if cue in source and isinstance(palette.get(key), dict)
            ),
            None,
        )
        palette = (
            palette[variant]
            if variant
            else next((value for value in palette.values() if isinstance(value, dict) and "accent" in value), {})
        )
    if not isinstance(palette, dict):
        palette = {}
    defaults = {"bg": "FFFFFF", "surface": "F4F8F1", "accent": "00561F", "text": "1E2A22", "muted": "617067"}
    return {key: str(palette.get(key, value)).lstrip("#") for key, value in defaults.items()}


def rgb(value: str) -> RGBColor:
    return RGBColor.from_string(value)


def font_name(spec: dict[str, Any]) -> str:
    fonts = spec.get("fonts", {})
    def first_font(value: Any) -> str | None:
        if isinstance(value, list):
            return first_font(value[0]) if value else None
        if isinstance(value, dict):
            return first_font(value.get("font") or value.get("name") or value.get("primary"))
        return value if isinstance(value, str) and value.strip() else None

    if isinstance(fonts, dict):
        return first_font(fonts.get("primary") or fonts.get("heading") or fonts.get("fallback")) or "微软雅黑"
    return first_font(fonts) or "微软雅黑"


def add_text(slide, text: str, left: float, top: float, width: float, height: float, size: float, color: str, font: str, bold: bool = False) -> None:
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    paragraph = box.text_frame.paragraphs[0]
    paragraph.text = text
    paragraph.font.name = font
    paragraph.font.size = Pt(size)
    paragraph.font.bold = bold
    paragraph.font.color.rgb = rgb(color)


def add_picture_contain(slide, asset_path: Path, left: float, top: float, width: float, height: float) -> None:
    with Image.open(asset_path) as image:
        ratio = image.width / max(image.height, 1)
    box_ratio = width / height
    if ratio >= box_ratio:
        picture_width = width
        picture_height = width / ratio
    else:
        picture_height = height
        picture_width = height * ratio
    x = left + (width - picture_width) / 2
    y = top + (height - picture_height) / 2
    slide.shapes.add_picture(
        str(asset_path),
        Inches(x),
        Inches(y),
        width=Inches(picture_width),
        height=Inches(picture_height),
    )


def make_slide(
    prs: Presentation,
    brief: dict[str, Any],
    spec: dict[str, Any],
    role: str,
    asset_path: Path | None,
) -> None:
    palette = palette_for(spec)
    font = font_name(spec)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    background = slide.background.fill
    background.solid()
    background.fore_color.rgb = rgb(palette["bg"])
    subtitle = str(brief.get("subtitle") or brief.get("purpose") or "")
    author = str(brief.get("author") or brief.get("audience") or "")
    date = str(brief.get("date") or "")
    footer = "  ".join(part for part in (author, date) if part)
    if role == "safe":
        band = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.18))
        band.fill.solid(); band.fill.fore_color.rgb = rgb(palette["accent"]); band.line.fill.background()
        accent = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.85), Inches(1.25), Inches(0.10), Inches(3.8))
        accent.fill.solid(); accent.fill.fore_color.rgb = rgb(palette["accent"]); accent.line.fill.background()
        add_text(slide, str(brief["title"]), 1.25, 1.32, 10.8, 1.35, 30, palette["text"], font, True)
        if subtitle:
            add_text(slide, subtitle, 1.27, 2.95, 9.8, 0.65, 18, palette["muted"], font)
        if footer:
            add_text(slide, footer, 1.27, 5.85, 8.5, 0.45, 13, palette["muted"], font)
        if asset_path:
            add_picture_contain(slide, asset_path, 9.65, 0.55, 2.45, 0.62)
    elif role == "structured":
        header = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(1.05))
        header.fill.solid(); header.fill.fore_color.rgb = rgb(palette["accent"]); header.line.fill.background()
        add_text(slide, str(brief["title"]), 0.85, 0.28, 11.55, 0.55, 25, palette["bg"], font, True)
        panel = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.85), Inches(1.62), Inches(11.63), Inches(2.15))
        panel.fill.solid(); panel.fill.fore_color.rgb = rgb(palette["surface"]); panel.line.color.rgb = rgb(palette["accent"])
        if subtitle:
            add_text(slide, subtitle, 1.25, 2.22, 10.8, 0.65, 19, palette["text"], font)
        if footer:
            add_text(slide, footer, 8.0, 6.22, 4.35, 0.38, 13, palette["muted"], font)
        if asset_path:
            add_picture_contain(slide, asset_path, 8.85, 4.25, 3.0, 1.15)
    else:
        block = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(7.95), Inches(0), Inches(5.383), Inches(7.5))
        block.fill.solid(); block.fill.fore_color.rgb = rgb(palette["surface"]); block.line.fill.background()
        rule = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.0), Inches(1.1), Inches(4.6), Inches(0.12))
        rule.fill.solid(); rule.fill.fore_color.rgb = rgb(palette["accent"]); rule.line.fill.background()
        add_text(slide, str(brief["title"]), 1.0, 1.55, 6.35, 2.05, 34, palette["text"], font, True)
        if subtitle:
            add_text(slide, subtitle, 1.02, 4.12, 5.95, 1.0, 18, palette["muted"], font)
        if footer:
            add_text(slide, footer, 8.55, 5.9, 3.75, 0.55, 13, palette["text"], font)
        if asset_path:
            add_picture_contain(slide, asset_path, 8.65, 1.35, 3.9, 2.0)


def pillow_font(size: int) -> ImageFont.ImageFont:
    for candidate in ("C:/Windows/Fonts/msyh.ttc", "C:/Windows/Fonts/simsun.ttc", "arial.ttf"):
        try:
            return ImageFont.truetype(candidate, size)
        except OSError:
            continue
    return ImageFont.load_default()


def paste_contain(image: Image.Image, asset_path: Path, box: tuple[int, int, int, int]) -> None:
    left, top, right, bottom = box
    width, height = right - left, bottom - top
    with Image.open(asset_path) as source:
        source = source.convert("RGBA")
        source.thumbnail((width, height), Image.Resampling.LANCZOS)
        x = left + (width - source.width) // 2
        y = top + (height - source.height) // 2
        image.paste(source, (x, y), source)


def write_png(
    path: Path,
    brief: dict[str, Any],
    spec: dict[str, Any],
    role: str,
    asset_path: Path | None,
) -> None:
    palette = palette_for(spec)
    image = Image.new("RGB", (1600, 900), f"#{palette['bg']}")
    draw = ImageDraw.Draw(image)
    subtitle = str(brief.get("subtitle") or brief.get("purpose") or "")
    footer = "  ".join(str(part) for part in (brief.get("author") or brief.get("audience") or "", brief.get("date") or "") if part)
    if role == "safe":
        draw.rectangle((0, 0, 1600, 22), fill=f"#{palette['accent']}")
        draw.rectangle((102, 150, 114, 610), fill=f"#{palette['accent']}")
        draw.text((150, 160), str(brief["title"]), font=pillow_font(50), fill=f"#{palette['text']}")
        if subtitle:
            draw.text((150, 355), subtitle, font=pillow_font(30), fill=f"#{palette['muted']}")
        if footer:
            draw.text((150, 700), footer, font=pillow_font(22), fill=f"#{palette['muted']}")
        if asset_path:
            paste_contain(image, asset_path, (1160, 60, 1450, 135))
    elif role == "structured":
        draw.rectangle((0, 0, 1600, 126), fill=f"#{palette['accent']}")
        draw.text((102, 34), str(brief["title"]), font=pillow_font(40), fill=f"#{palette['bg']}")
        draw.rectangle((102, 194, 1498, 452), fill=f"#{palette['surface']}", outline=f"#{palette['accent']}", width=2)
        if subtitle:
            draw.text((150, 270), subtitle, font=pillow_font(30), fill=f"#{palette['text']}")
        if footer:
            draw.text((960, 746), footer, font=pillow_font(22), fill=f"#{palette['muted']}")
        if asset_path:
            paste_contain(image, asset_path, (1060, 510, 1420, 650))
    else:
        draw.rectangle((954, 0, 1600, 900), fill=f"#{palette['surface']}")
        draw.rectangle((120, 132, 672, 144), fill=f"#{palette['accent']}")
        draw.text((120, 190), str(brief["title"]), font=pillow_font(56), fill=f"#{palette['text']}")
        if subtitle:
            draw.text((122, 495), subtitle, font=pillow_font(30), fill=f"#{palette['muted']}")
        if footer:
            draw.text((1025, 708), footer, font=pillow_font(22), fill=f"#{palette['text']}")
        if asset_path:
            paste_contain(image, asset_path, (1040, 165, 1480, 405))
    image.save(path)


def sha256_file(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def normalize_pptx(path: Path) -> None:
    normalized = path.with_suffix(".normalized.pptx")
    with zipfile.ZipFile(path, "r") as source, zipfile.ZipFile(normalized, "w") as target:
        for original in source.infolist():
            info = zipfile.ZipInfo(original.filename, (1980, 1, 1, 0, 0, 0))
            info.compress_type = original.compress_type
            info.external_attr = original.external_attr
            info.create_system = original.create_system
            target.writestr(info, source.read(original.filename))
    os.replace(normalized, path)


def render_with_powerpoint(pptx_path: Path, output_dir: Path) -> None:
    script = Path(__file__).with_name("export_style_discovery_previews.ps1")
    result = subprocess.run(
        [
            "powershell",
            "-NoProfile",
            "-ExecutionPolicy",
            "Bypass",
            "-File",
            str(script),
            "-Pptx",
            str(pptx_path),
            "-OutputDir",
            str(output_dir),
        ],
        text=True,
        capture_output=True,
        encoding="utf-8",
    )
    if result.returncode != 0:
        raise RuntimeError((result.stdout + result.stderr).strip() or "PowerPoint export failed")


def write_contact_sheet(path: Path, preview_paths: list[Path], options: list[dict[str, Any]]) -> None:
    sheet = Image.new("RGB", (1600, 620), "#FFFFFF")
    draw = ImageDraw.Draw(sheet)
    draw.text((50, 28), "同一科研题目的三种真实视觉方向", font=pillow_font(30), fill="#1E2A22")
    role_labels = {
        "safe": "SAFE / 官方保真",
        "structured": "STRUCTURED / 学术结构",
        "exploratory": "EXPLORATORY / 候选方向",
    }
    for index, preview_path in enumerate(preview_paths):
        with Image.open(preview_path) as source:
            preview = source.resize((480, 270))
        x = 50 + index * 535
        sheet.paste(preview, (x, 100))
        option = options[index]
        draw.text(
            (x, 398),
            f"{option['option']}  {role_labels[option['role']]}",
            font=pillow_font(23),
            fill="#1E2A22",
        )
        draw.text((x, 445), option["name"], font=pillow_font(18), fill="#617067")
    sheet.save(path)


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--brief", type=Path, required=True)
    parser.add_argument("--asset-review", type=Path, required=True)
    parser.add_argument("--output-dir", type=Path, required=True)
    parser.add_argument("--select", choices=OPTIONS)
    parser.add_argument("--render-method", choices=("auto", "powerpoint", "pillow"), default="auto")
    parser.add_argument("--showcase", action="store_true", help="Also write style-discovery-showcase.pptx")
    args = parser.parse_args()
    try:
        brief = load_json(args.brief)
        errors = validate_brief(brief)
        asset_review = load_json(args.asset_review)
        asset_errors, approved_assets = validate_asset_review(asset_review, args.asset_review)
        errors.extend(asset_errors)
        if errors:
            raise ValueError("; ".join(errors))
        selected = select_styles(brief)
    except (OSError, ValueError, json.JSONDecodeError) as error:
        print(f"ERROR: {error}")
        return 1

    args.output_dir.mkdir(parents=True, exist_ok=True)
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    prs.core_properties.title = str(brief["title"])
    prs.core_properties.subject = "AI_PPT Visual Discovery"
    prs.core_properties.author = "AI_PPT"
    prs.core_properties.last_modified_by = "AI_PPT"
    prs.core_properties.created = FIXED_TIMESTAMP
    prs.core_properties.modified = FIXED_TIMESTAMP
    prs.core_properties.revision = 1
    preview_paths: list[Path] = []
    options: list[dict[str, Any]] = []
    preview_specs: list[tuple[str, dict[str, Any]]] = []
    preview_asset = approved_assets[0] if approved_assets else None
    for letter, (role, (entry, spec, _score)) in zip(OPTIONS, selected, strict=True):
        make_slide(prs, brief, spec, role, preview_asset)
        preview_path = args.output_dir / f"option-{letter.lower()}.png"
        preview_paths.append(preview_path)
        preview_specs.append((role, spec))
        options.append(
            {
                "option": letter,
                "role": role,
                "style_id": entry["id"],
                "name": entry.get("name", entry["id"]),
                "generation_status": entry.get("generation_status"),
            }
        )
    style_options_path = args.output_dir / "style-options.pptx"
    prs.save(style_options_path)
    normalize_pptx(style_options_path)
    if args.showcase:
        shutil.copyfile(style_options_path, args.output_dir / "style-discovery-showcase.pptx")
    render_method = args.render_method
    if render_method in {"auto", "powerpoint"}:
        try:
            render_with_powerpoint(style_options_path.resolve(), args.output_dir.resolve())
            render_method = "powerpoint"
        except RuntimeError as error:
            if args.render_method == "powerpoint":
                print(f"ERROR: {error}")
                return 1
            print(f"WARN: PowerPoint rendering unavailable; using Pillow fallback: {error}")
            render_method = "pillow"
    if render_method == "pillow":
        for preview_path, (role, spec) in zip(preview_paths, preview_specs, strict=True):
            write_png(preview_path, brief, spec, role, preview_asset)
    write_contact_sheet(args.output_dir / "contact-sheet.png", preview_paths, options)
    chosen = next((item for item in options if item["option"] == args.select), None)
    selection = {
        "schema_version": 1,
        "generator_version": GENERATOR_VERSION,
        "runtime": {
            "python": sys.version.split()[0],
            "python_pptx": version("python-pptx"),
            "pillow": version("Pillow"),
        },
        "options": options,
        "selected_option": args.select,
        "selected_style_id": chosen["style_id"] if chosen else None,
        "confirmed": bool(args.select),
        "approval_scope": "deck_only" if chosen and chosen["generation_status"] == "style_selection_only" else "none",
        "render_method": render_method,
        "approved_asset_count": len(approved_assets),
        "input_hashes": {
            "brief_sha256": sha256_file(args.brief),
            "asset_review_sha256": sha256_file(args.asset_review),
            "style_registry_sha256": sha256_file(REGISTRY_PATH),
            "generator_sha256": sha256_file(Path(__file__)),
            "style_specs_sha256": {
                entry["id"]: sha256_file(ROOT / entry["style_spec"])
                for _role, (entry, _spec, _score) in selected
            },
        },
    }
    (args.output_dir / "style-selection.json").write_text(
        json.dumps(selection, ensure_ascii=False, indent=2, sort_keys=True) + "\n",
        encoding="utf-8",
    )
    print(f"Generated style discovery outputs in {args.output_dir}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
